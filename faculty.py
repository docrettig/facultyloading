"""Roster validation, schedule matching, and downloadable unsent email drafts."""
from io import BytesIO
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import getaddresses
import re
import unicodedata
from openpyxl import load_workbook
from loading import load_limits


def normalize_name(name):
    return ' '.join(unicodedata.normalize('NFKC',str(name)).casefold().split())


def addresses(value, label):
    value=str(value or '').strip()
    if '\r' in value or '\n' in value:raise ValueError(f'Invalid {label} email address.')
    found=getaddresses([value.replace(';',',')])
    emails=[a for _,a in found]
    if not emails or any(not re.fullmatch(r'[^\s@,;<>]+@[^\s@,;<>]+\.[^\s@,;<>]+',a) for a in emails):
        raise ValueError(f'Missing or invalid {label} email address.')
    return emails


def read_roster(blob, term):
    book=load_workbook(BytesIO(blob),read_only=True,data_only=True)
    sheet=book.active
    rows=list(sheet.iter_rows(values_only=False));entries=[];errors=[];seen=set()
    if not rows or [str(c.value or '').strip() for c in rows[0][:6]]!=['Instructor','Loading %','Role','Chair','Admin','VCAA']:
        raise ValueError('Faculty roster must have columns Instructor, Loading %, Role, Chair, Admin, VCAA in A–F.')
    for row in rows[1:]:
        if all(c.value is None for c in row[:6]):continue
        name=str(row[0].value or '').strip()
        try:
            if not name:raise ValueError('Missing faculty name.')
            key=normalize_name(name)
            if key in seen:raise ValueError('Duplicate faculty name in roster.')
            seen.add(key)
            cell=row[1];value=cell.value
            if isinstance(value,str):
                request=float(value.strip().rstrip('%'))/100
            elif isinstance(value,(int,float)) and not isinstance(value,bool):
                request=float(value) if '%' in cell.number_format else float(value)/100
            else:raise ValueError('Missing loading percentage.')
            if not 0<=request<=1:raise ValueError('Loading percentage must be between 0 and 100%.')
            role=str(row[2].value or '').strip();load_limits(term,role)
            entries.append(dict(name=name,loading=request,role=role,to=addresses(row[3].value,'chair'),cc=addresses(row[4].value,'admin assistant')+addresses(row[5].value,'VCAA')))
        except (ValueError,TypeError) as error:errors.append(f'{name or "Roster row "+str(row[0].row)}: {error}')
    book.close()
    return entries,errors


def match_roster(entries,people):
    index={}
    for person,rows in people.items():index.setdefault(normalize_name(person[1]),[]).append((person,rows))
    matched=[];errors=[]
    for entry in entries:
        candidates=index.get(normalize_name(entry['name']),[])
        if not candidates:errors.append(f'No match for {entry["name"]}.');continue
        if len(candidates)>1:errors.append(f'Multiple schedule matches for {entry["name"]}; review instructor IDs.');continue
        person,rows=candidates[0];matched.append((entry,person,rows))
    return matched,errors


def email_draft(entry,term,filename,workbook):
    msg=EmailMessage(policy=SMTP)
    msg['To']=', '.join(entry['to'])
    msg['Cc']=', '.join(dict.fromkeys(entry['cc']))
    msg['Subject']=f'Loading document for review: {entry["name"]} – {term}'
    msg['X-Unsent']='1'
    msg.set_content(f'Please review the attached loading document for {entry["name"]} for {term}.\n\nRole: {entry["role"]}\nLoad request: {entry["loading"]:.0%}\n\nPlease confirm the assignments and proposed overloads, or reply with corrections.\n')
    msg.add_attachment(workbook,maintype='application',subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename=filename)
    return msg.as_bytes()
